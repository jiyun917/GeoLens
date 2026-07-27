"""Fill missing ± sample SD values into the per-model tables of the paper's
narrative doc GeoLens_최종결과_통합최종본.docx.

Reads recomputed cell values from table1_full_sd.md (built by
scripts/compute_table1_sd.py). Preserves editorial annotations (e.g. the
'*' outlier marker on Qwen vanilla cost) and 'n/a' cells for Qwen
long_context. Writes back in place; makes a .bak backup first.

Idempotent: cells that already have '±' are re-formatted from the same
source, so re-running does not double-append.
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import docx  # python-docx

ROOT = Path(__file__).resolve().parent.parent
DOCX_PATH = Path(r"C:\Users\user\Downloads\GeoLens_최종결과_통합최종본.docx")
TABLE1_MD = ROOT / "data" / "eval" / "results" / "table1_full_sd.md"

# Table index in the docx → model tag used in the md
DOCX_TABLE_TO_MODEL = {
    1: "gemini",
    2: "claude",
    3: "gpt",
    4: "qwen",
}

# Docx metric-row label (first col text) → md metric label
METRIC_ROW_MAP = {
    "Step Acc":  "Step Acc",
    "Hall Rate": "Hall Rate",
    "Loop Rate": "Loop Rate",
    "Goal Comp": "Goal Comp",
    "Latency":   "Latency",
    "Cost $m":   "Cost/step",
}

# Docx backend column labels → md backend labels
BACKEND_COL_MAP = {
    "no_rag":       "no_rag",
    "vanilla":      "vanilla_vector",
    "graph_only":   "graph_only",
    "vision_only":  "vision_only",
    "full_system":  "full_system",
    "state_path":   "state_path",
    "long_context": "long_context",
}


def parse_table1_md(path: Path) -> dict:
    """Returns {(model, backend, metric_label) : "mean±sd" display string}
    from table1_full_sd.md (both core and goal-comp tables).
    """
    text = path.read_text(encoding="utf-8")
    out: dict = {}
    lines = text.splitlines()
    header: list[str] | None = None
    metric_labels: list[str] | None = None
    for line in lines:
        s = line.strip()
        if not s.startswith("|"):
            header = None
            metric_labels = None
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if header is None:
            header = cells
            # Extract metric label from each header cell like "Step Acc (%) ↑"
            metric_labels = []
            for c in cells[1:]:
                # take the part before "("
                base = c.split("(")[0].strip()
                metric_labels.append(base)
            continue
        if all(c.startswith("---") or c == "" for c in cells):
            continue
        # Data row: "gemini / no_rag" | v1 | v2 | ...
        label = cells[0]
        if " / " not in label:
            continue
        model, backend = [p.strip() for p in label.split("/", 1)]
        for i, val in enumerate(cells[1:]):
            metric_lbl = metric_labels[i] if metric_labels and i < len(metric_labels) else None
            if metric_lbl is None:
                continue
            out[(model, backend, metric_lbl)] = val
    return out


COST_RE = re.compile(r"^\$(\d+\.?\d*)±(\d+\.?\d*)m$")


def format_for_docx(display: str, metric_label: str) -> str:
    """Convert md-style display ('$27.13±0.05m') to docx-style ('27.13±0.05').

    Cost cells drop the leading '$' and trailing 'm' because the column is
    already labeled "Cost $m" in the docx. Other cells pass through.
    """
    if metric_label == "Cost/step":
        m = COST_RE.match(display)
        if m:
            return f"{m.group(1)}±{m.group(2)}"
    return display


def existing_annotation(old: str) -> str:
    """Extract trailing non-numeric annotation from cells like '60.45*' → '*'."""
    m = re.search(r"([^\d\s±.]+)\s*$", old.strip())
    if m and m.group(1) not in ("m", "s", "%"):
        return m.group(1)
    return ""


def update_cell(cell, new_text: str) -> None:
    """Rewrite a docx cell's text while preserving the first paragraph's
    style. Overwrites all runs in the first paragraph and drops others.
    """
    if not cell.paragraphs:
        cell.text = new_text
        return
    para = cell.paragraphs[0]
    # Remove extra paragraphs (rare in these tables).
    for extra in list(cell.paragraphs[1:]):
        p_el = extra._element
        p_el.getparent().remove(p_el)
    # Clear existing runs; keep first run's formatting if present.
    if para.runs:
        first = para.runs[0]
        first.text = new_text
        for r in para.runs[1:]:
            r_el = r._element
            r_el.getparent().remove(r_el)
    else:
        para.add_run(new_text)


def main():
    if not DOCX_PATH.exists():
        print(f"[fatal] docx not found: {DOCX_PATH}", file=sys.stderr)
        sys.exit(1)
    if not TABLE1_MD.exists():
        print(f"[fatal] md not found: {TABLE1_MD} — run compute_table1_sd.py first",
              file=sys.stderr)
        sys.exit(1)

    values = parse_table1_md(TABLE1_MD)
    print(f"parsed {len(values)} cell values from {TABLE1_MD.name}")

    backup = DOCX_PATH.with_suffix(DOCX_PATH.suffix + ".bak")
    if not backup.exists():
        shutil.copy2(DOCX_PATH, backup)
        print(f"backup created: {backup.name}")
    else:
        print(f"backup already present: {backup.name} (kept)")

    doc = docx.Document(str(DOCX_PATH))
    changed = 0
    kept_na = 0
    for tbl_idx, model in DOCX_TABLE_TO_MODEL.items():
        table = doc.tables[tbl_idx]
        # Header row: first row, backend labels in cols 1..N.
        header_row = table.rows[0]
        backend_by_col: dict[int, str] = {}
        for ci, c in enumerate(header_row.cells):
            txt = c.text.strip()
            for docx_lbl, md_lbl in BACKEND_COL_MAP.items():
                if txt == docx_lbl:
                    backend_by_col[ci] = md_lbl
                    break

        # Data rows
        for row in table.rows[1:]:
            metric_cell = row.cells[0].text.strip()
            # Match to metric label using startswith (docx has "Step Acc ↑" etc.)
            md_metric = None
            for docx_lbl, md_lbl in METRIC_ROW_MAP.items():
                if metric_cell.startswith(docx_lbl):
                    md_metric = md_lbl
                    break
            if md_metric is None:
                continue
            for col_idx, backend in backend_by_col.items():
                cell = row.cells[col_idx]
                old = cell.text.strip()
                if old == "n/a" or old == "":
                    kept_na += 1
                    continue
                key = (model, backend, md_metric)
                new_display = values.get(key)
                if new_display is None:
                    print(f"[warn] no value for {model}/{backend}/{md_metric}",
                          file=sys.stderr)
                    continue
                annotation = existing_annotation(old)
                formatted = format_for_docx(new_display, md_metric)
                if annotation:
                    formatted = formatted + annotation
                if old != formatted:
                    update_cell(cell, formatted)
                    changed += 1

    doc.save(str(DOCX_PATH))
    print(f"cells updated: {changed}")
    print(f"cells kept (n/a or empty): {kept_na}")


if __name__ == "__main__":
    main()
